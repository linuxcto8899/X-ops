import io
from django.utils import timezone
from rest_framework import viewsets, status
from rest_framework.decorators import action
from rest_framework.response import Response
from openpyxl import load_workbook
from .models import Host, HostGroup, Credential
from .serializers import HostSerializer, HostGroupSerializer, CredentialSerializer
from .ssh import check_host


class HostGroupViewSet(viewsets.ModelViewSet):
    queryset = HostGroup.objects.all()
    serializer_class = HostGroupSerializer
    search_fields = ['name']


class CredentialViewSet(viewsets.ModelViewSet):
    queryset = Credential.objects.all()
    serializer_class = CredentialSerializer
    search_fields = ['name', 'username']


class HostViewSet(viewsets.ModelViewSet):
    queryset = Host.objects.all()
    serializer_class = HostSerializer
    search_fields = ['name', 'hostname', 'username']
    filterset_fields = ['status']

    def get_queryset(self):
        qs = super().get_queryset()
        group = self.request.query_params.get('group')
        if group:
            qs = qs.filter(groups__id=group)
        return qs.prefetch_related('groups')

    @action(detail=True, methods=['post'])
    def check(self, request, pk=None):
        host = self.get_object()
        ok, detail = check_host(host)
        host.last_checked_at = timezone.now()
        host.save(update_fields=['last_checked_at'])
        return Response({'ok': ok, 'detail': detail, 'status': host.status})

    @action(detail=False, methods=['post'])
    def batch_check(self, request):
        ids = request.data.get('ids', [])
        hosts = Host.objects.filter(id__in=ids)
        results = []
        for host in hosts:
            ok, detail = check_host(host)
            host.last_checked_at = timezone.now()
            host.save(update_fields=['last_checked_at'])
            results.append({'id': host.id, 'name': host.name, 'ok': ok,
                            'detail': detail, 'status': host.status})
        return Response(results)

    @action(detail=False, methods=['post'])
    def import_excel(self, request):
        """Import hosts from an uploaded .xlsx file.

        Expected columns: name, hostname, port, username, password, group
        """
        file = request.FILES.get('file')
        if not file:
            return Response({'detail': 'file is required'}, status=status.HTTP_400_BAD_REQUEST)
        try:
            wb = load_workbook(io.BytesIO(file.read()), read_only=True)
            ws = wb.active
            rows = list(ws.iter_rows(values_only=True))
        except Exception as exc:
            return Response({'detail': f'Invalid Excel: {exc}'}, status=status.HTTP_400_BAD_REQUEST)
        if not rows:
            return Response({'detail': 'Empty file'}, status=status.HTTP_400_BAD_REQUEST)
        header = [str(c).strip().lower() if c else '' for c in rows[0]]
        created, failed = [], []
        for row in rows[1:]:
            if not row or not row[0]:
                continue
            data = dict(zip(header, row))
            try:
                host = Host.objects.create(
                    name=str(data.get('name', '')).strip(),
                    hostname=str(data.get('hostname', '')).strip(),
                    port=int(data.get('port') or 22),
                    username=str(data.get('username', '')).strip(),
                )
                host.set_password(str(data.get('password') or ''))
                host.save()
                group_name = str(data.get('group') or '').strip()
                if group_name:
                    group, _ = HostGroup.objects.get_or_create(name=group_name)
                    host.groups.add(group)
                created.append({'id': host.id, 'name': host.name})
            except Exception as exc:
                failed.append({'row': str(row[0]), 'error': str(exc)})
        return Response({'created': created, 'failed': failed})
